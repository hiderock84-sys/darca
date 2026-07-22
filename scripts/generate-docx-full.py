#!/usr/bin/env python3
"""PDF（A4版）の各ページを画像化し、見た目そのままの A4 Word 文書を作る。

各ページを1枚の画像として A4・余白ゼロで配置するため、PDF とまったく同じ
レイアウト・写真・図解・配色を Word 上で再現できる（＝「全てそのまま」）。

使い方:
    python3 scripts/generate-docx-full.py public/manual-a4.pdf public/manual.docx
"""
import io
import sys
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_BREAK

src = sys.argv[1] if len(sys.argv) > 1 else "public/manual-a4.pdf"
out = sys.argv[2] if len(sys.argv) > 2 else "public/manual.docx"

pdf = fitz.open(src)

doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(0)
sec.bottom_margin = Mm(0)
sec.left_margin = Mm(0)
sec.right_margin = Mm(0)
sec.header_distance = Mm(0)
sec.footer_distance = Mm(0)

# 各ページを高解像度で画像化して全面に貼り付ける
zoom = 2.0  # ≒144dpi（品質とファイルサイズのバランス）
for i in range(pdf.page_count):
    pix = pdf[i].get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    img = io.BytesIO(pix.tobytes("jpeg", jpg_quality=82))

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    run = p.add_run()
    # A4 全面に配置（余白ゼロ）。高さは 296.8mm にして端数による送りを防ぐ
    run.add_picture(img, width=Mm(210), height=Mm(296.8))

    if i < pdf.page_count - 1:
        br = doc.add_paragraph()
        br.paragraph_format.space_before = Pt(0)
        br.paragraph_format.space_after = Pt(0)
        br.paragraph_format.line_spacing = Pt(1)
        br.add_run().add_break(WD_BREAK.PAGE)

doc.save(out)
print(f"DOCX (faithful) generated: {out}  ({pdf.page_count} pages)")
