#!/usr/bin/env python3
"""編集可能で完成版と同等クオリティの Word を生成する。

src/data/manual.ts をダンプした JSON（scripts/manual.json）を読み、
章見出しの色帯・POINTカード・囲み(コラム)・Q&A・比較表・写真などを
ネイティブ Word 要素（表・セル背景色・スタイル・埋め込み画像）で再現する。

前提: npx tsx scripts/dump-manual.ts scripts/manual.json を先に実行。
使い方: python3 scripts/generate-docx-rich.py scripts/manual.json public/manual.docx
"""
import io
import json
import os
import sys

from PIL import Image  # webp→jpg 変換に使用
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else "scripts/manual.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else "public/manual.docx"
IMG_DIR = "src/assets/manual"

NAVY = RGBColor(0x1B, 0x33, 0x58)
NAVY_HEX = "1B3358"
NAVY_DEEP_HEX = "12233F"
GREEN = RGBColor(0x2F, 0x9E, 0x6B)
GREEN_HEX = "2F9E6B"
WARM = RGBColor(0xB0, 0x6A, 0x2C)
BLUE = RGBColor(0x3A, 0x86, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2B, 0x33, 0x42)
GRAY = RGBColor(0x60, 0x6A, 0x7A)

TINT_GRAY = "F3F5F9"
TINT_GREEN = "E9F4EE"
TINT_BLUE = "EAF1FA"
TINT_RED = "FBEDED"
TINT_WARM = "F7F0E4"

CONTENT_W = 186  # mm（A4 - 左右余白12mm）

data = json.load(open(JSON_PATH, encoding="utf-8"))

doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
for m in ("top_margin", "bottom_margin"):
    setattr(sec, m, Mm(12))
sec.left_margin = Mm(12)
sec.right_margin = Mm(12)

# 既定フォント（日本語含む）
normal = doc.styles["Normal"]
normal.font.name = "Yu Gothic"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
rfonts.set(qn("w:eastAsia"), "Yu Gothic")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(4)


def _shade(el, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    el.append(shd)


def cell_bg(cell, hex_color):
    _shade(cell._tc.get_or_add_tcPr(), hex_color)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for name, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def run(p, text, *, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    r.font.name = "Yu Gothic"
    r.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Yu Gothic")
    return r


def para(text=None, *, style=None, space_before=0, space_after=4, align=None, **kw):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text is not None:
        run(p, text, **kw)
    return p


def spacer(pt=4):
    para(space_after=pt)


def img_stream(name, max_w=1600):
    """webp などを読み込み JPEG バイト列で返す。失敗時 None。"""
    for ext in (".webp", ".jpg", ".png"):
        path = os.path.join(IMG_DIR, name + ext)
        if os.path.exists(path):
            try:
                im = Image.open(path).convert("RGB")
                if im.width > max_w:
                    h = round(im.height * max_w / im.width)
                    im = im.resize((max_w, h), Image.LANCZOS)
                buf = io.BytesIO()
                im.save(buf, format="JPEG", quality=84)
                buf.seek(0)
                return buf
            except Exception as e:
                sys.stderr.write(f"img fail {name}: {e}\n")
                return None
    return None


def add_image(name, caption=None, width=CONTENT_W):
    st = img_stream(name)
    if not st:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(st, width=Mm(width))
    if caption:
        c = para(space_after=8)
        run(c, caption, size=8.5, color=GRAY, italic=True)


def one_cell_box(fill_hex):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Mm(CONTENT_W)
    cell = t.rows[0].cells[0]
    cell.width = Mm(CONTENT_W)
    no_borders(t)
    cell_bg(cell, fill_hex)
    cell_margins(cell)
    # 最初の空段落を再利用
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    return cell


def cell_para(cell, text=None, *, first=False, space_after=2, **kw):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.45
    if text is not None:
        run(p, text, **kw)
    return p


# ---------- 見出し系 ----------
def chapter_band(no, title, catch):
    cell = one_cell_box(NAVY_HEX)
    cell_margins(cell, top=150, bottom=150, left=170, right=170)
    cell_para(cell, no, first=True, size=9, bold=True, color=WHITE, space_after=3)
    cell_para(cell, title, size=19, bold=True, color=WHITE, space_after=4)
    if catch:
        cell_para(cell, catch, size=11.5, bold=True, color=RGBColor(0xF2, 0xE7, 0xD6))
    spacer(6)


def h_section(text):
    p = para(space_before=8, space_after=3)
    run(p, text, bold=True, size=13, color=NAVY)


def lead(lines):
    for t in lines:
        p = para(space_before=2, space_after=6)
        p.paragraph_format.left_indent = Mm(3)
        run(p, t, size=11, bold=True, color=NAVY)


def body_lines(lines):
    for t in lines:
        para(t)


def quote(text):
    p = para(space_before=4, space_after=6)
    p.paragraph_format.left_indent = Mm(4)
    run(p, text, bold=True, italic=True, color=NAVY, size=11.5)


# ---------- 部品 ----------
def caption_line(text):
    p = para(space_before=1, space_after=6)
    run(p, text, size=8.5, color=GRAY)


def diagram_title(text):
    p = para(space_before=6, space_after=3)
    run(p, "◆ " + text, bold=True, size=11.5, color=NAVY)


def cards(items, cols=3, fill=TINT_GRAY, label=None, label_color=BLUE):
    n = len(items)
    cols = min(cols, n)
    rows = (n + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    cw = Mm(CONTENT_W / cols)
    for ci in range(cols):
        t.columns[ci].width = cw
    for idx, it in enumerate(items):
        r, c = divmod(idx, cols)
        cell = t.rows[r].cells[c]
        cell.width = cw
        cell_bg(cell, fill)
        cell_margins(cell, top=90, bottom=90, left=120, right=120)
        if label:
            cell_para(cell, label, first=True, size=7.5, bold=True, color=label_color, space_after=2)
            cell_para(cell, it["title"], size=10.5, bold=True, color=NAVY, space_after=2)
        else:
            cell_para(cell, it["title"], first=True, size=10.5, bold=True, color=NAVY, space_after=2)
        cell_para(cell, it.get("body") or it.get("note") or "", size=9, space_after=0)
    spacer(6)


def bullets(items, mark="・", color=INK):
    for it in items:
        p = para(space_after=2)
        run(p, mark + " ", color=color, bold=True)
        run(p, it)


def checklist(items):
    for it in items:
        p = para(space_after=2)
        run(p, "☐ ", bold=True, color=NAVY)
        run(p, it)


def color_box(label, title, lines, fill, accent):
    cell = one_cell_box(fill)
    if label:
        cell_para(cell, label, first=True, size=8, bold=True, color=accent, space_after=2)
        if title:
            cell_para(cell, title, size=11.5, bold=True, color=NAVY, space_after=3)
    elif title:
        cell_para(cell, title, first=True, size=11.5, bold=True, color=NAVY, space_after=3)
    for i, t in enumerate(lines):
        cell_para(cell, t, size=10, space_after=2)
    spacer(6)


def callout(text, fill=TINT_WARM, accent=WARM, label=None):
    cell = one_cell_box(fill)
    if label:
        cell_para(cell, label, first=True, size=9.5, bold=True, color=accent, space_after=2)
        cell_para(cell, text, size=10, space_after=0)
    else:
        cell_para(cell, text, first=True, size=10, space_after=0)
    spacer(6)


def steps(items):
    for i, s in enumerate(items, 1):
        p = para(space_after=3)
        run(p, f"{i}. ", bold=True, size=11, color=GREEN)
        run(p, s["title"], bold=True, color=NAVY)
        if s.get("note"):
            run(p, "  " + s["note"], size=9.5, color=GRAY)


def two_col(left, right):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    cw = Mm(CONTENT_W / 2)
    for side, cell, tint, accent in (
        (left, t.rows[0].cells[0], TINT_RED, RGBColor(0xC0, 0x5C, 0x5C)),
        (right, t.rows[0].cells[1], TINT_GREEN, GREEN),
    ):
        cell.width = cw
        t.columns[0 if side is left else 1].width = cw
        cell_bg(cell, tint)
        cell_margins(cell, top=90, bottom=90, left=120, right=120)
        cell_para(cell, side["title"], first=True, size=10.5, bold=True, color=accent, space_after=3)
        for it in side["items"]:
            pp = cell.add_paragraph()
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.line_spacing = 1.4
            run(pp, "・ ", bold=True, color=accent)
            run(pp, it, size=9.5)
    spacer(6)


def kv_table(rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    lw, rw = Mm(38), Mm(CONTENT_W - 38)
    for i, row in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.width = lw
        c1.width = rw
        cell_bg(c0, TINT_BLUE)
        cell_margins(c0, top=70, bottom=70)
        cell_margins(c1, top=70, bottom=70)
        cell_para(c0, row["label"], first=True, size=9.5, bold=True, color=NAVY, space_after=0)
        cell_para(c1, row["value"], first=True, size=9.5, space_after=0)
    spacer(6)


def qa(q, a):
    p = para(space_before=4, space_after=1)
    run(p, "Q ", bold=True, size=11, color=NAVY)
    run(p, q, bold=True, color=NAVY)
    p2 = para(space_after=6)
    run(p2, "A ", bold=True, size=11, color=GREEN)
    run(p2, a)


def case_box(c):
    tint = TINT_GREEN if c.get("tone") == "do" else TINT_RED
    accent = GREEN if c.get("tone") == "do" else RGBColor(0xC0, 0x5C, 0x5C)
    cell = one_cell_box(tint)
    cell_para(cell, f"{c['label']}　{c.get('person','')}", first=True, size=10.5, bold=True, color=accent, space_after=3)
    for t in c["story"]:
        cell_para(cell, t, size=9.5, space_after=2)
    if c.get("result"):
        cell_para(cell, "結果：" + c["result"], size=9.5, bold=True, color=NAVY, space_after=0)
    spacer(6)


# =====================================================================
# 本体
# =====================================================================
cover = data["cover"]
add_image(data["images"]["cover"], width=CONTENT_W)
p = para(space_before=6, space_after=2)
run(p, cover["category"], size=10, bold=True, color=BLUE)
p = para(space_after=4)
run(p, cover["title"].replace("\n", " "), size=26, bold=True, color=NAVY)
p = para(space_after=4)
run(p, cover["subtitle"], size=11, color=INK)
p = para(space_after=2)
run(p, cover["supply"], size=8.5, color=GRAY)
p = para(space_after=0)
run(p, data["org"]["name"], size=12, bold=True, color=NAVY)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 導入ストーリー
st = data["openingStory"]
chapter_band(st["chapterLabel"], st["title"], "")
add_image(data["images"]["story"], caption="夜、玄関にともる灯り。")
quote(st["lead"])
for t in st["paragraphs"]:
    quote(t) if t.startswith("「") else para(t)
color_box(None, None, st["closing"], TINT_BLUE, BLUE)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# はじめに
pf = data["preface"]
chapter_band("はじめに", pf["title"], "")
body_lines(pf["paragraphs"])
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 各章
chImages = data["chapterImages"]
for ch in data["chapters"]:
    chapter_band(ch["no"], ch["title"], ch.get("catch", ""))
    imgname = chImages.get(ch["id"])
    if imgname:
        add_image(imgname)
    lead(ch.get("lead", []))

    for key in ch:
        if key in ("no", "id", "title", "catch", "lead"):
            continue
        v = ch[key]
        if key == "sections":
            for s in v:
                h_section(s["heading"])
                body_lines(s["body"])
        elif key == "points":
            cards(v, cols=3, fill=TINT_GRAY, label="POINT")
        elif key == "cycle" or key == "flow":
            diagram_title(v["title"])
            steps(v["steps"])
            caption_line(v.get("caption", ""))
        elif key == "column" or key == "boundary" or key == "craft" or key == "consult":
            color_box(v.get("label", ""), v.get("title", ""), v.get("body", []), TINT_GREEN, GREEN)
        elif key == "checklist":
            diagram_title(v["title"])
            checklist(v["items"])
            caption_line(v.get("caption", ""))
        elif key == "examples":
            diagram_title(v["title"])
            for it in v["items"]:
                p = para(space_after=2)
                run(p, "・ " + it["do"], size=9.5, color=RGBColor(0xC0, 0x5C, 0x5C))
                run(p, "  →  " + it["instead"], size=9.5, bold=True, color=GREEN)
            caption_line(v.get("caption", ""))
        elif key == "patterns":
            diagram_title(v["title"])
            for it in v["items"]:
                p = para(space_after=2)
                run(p, "・ " + it["title"] + "：", bold=True, color=NAVY)
                run(p, it["note"], size=9.5)
            caption_line(v.get("caption", ""))
        elif key == "compare":
            diagram_title(v["title"])
            two_col(v["left"], v["right"])
            caption_line(v.get("caption", ""))
        elif key == "saylist":
            diagram_title(v["title"])
            two_col(
                {"title": v["bad"]["title"], "items": v["bad"]["items"]},
                {"title": v["good"]["title"], "items": v["good"]["items"]},
            )
        elif key == "safety":
            callout(v["body"], fill=TINT_RED, accent=RGBColor(0xC0, 0x5C, 0x5C), label="⚠ " + v["label"])
        elif key == "cases":
            for c in v:
                color_box(c["label"], c.get("from", ""), c["body"], TINT_GRAY, NAVY)
        elif key == "impersonation":
            callout(v["body"], fill=TINT_WARM, accent=WARM, label=v["label"])
        elif key == "note" or key == "reasonNote":
            callout(v, fill=TINT_WARM, accent=WARM)
        elif key == "dont":
            diagram_title(v["title"])
            for it in v["items"]:
                p = para(space_after=2)
                run(p, "✕ " + it["title"] + "：", bold=True, color=RGBColor(0xC0, 0x5C, 0x5C))
                run(p, it["note"], size=9.5)
        elif key == "do":
            diagram_title(v["title"])
            for it in v["items"]:
                p = para(space_after=2)
                run(p, "✓ " + it["title"] + "：", bold=True, color=GREEN)
                run(p, it["note"], size=9.5)
        elif key == "recovered" or key == "relapsed":
            case_box(v)
        elif key == "lesson":
            color_box(None, v["title"], [v["body"]], TINT_BLUE, BLUE)
        elif key == "groups":
            for g in v:
                h_section(g["label"])
                for it in g["items"]:
                    qa(it["q"], it["a"])
        elif key == "benefits":
            cards([{"title": b["title"], "body": b["body"]} for b in v], cols=3, fill=TINT_GREEN, label=None)
        elif key == "voices":
            diagram_title(v["title"])
            bullets([x.strip("「」") for x in v["items"]], mark="“", color=BLUE)
            caption_line(v.get("caption", ""))
        elif key == "program":
            diagram_title(v["title"])
            cards([{"title": it["title"], "body": it["body"]} for it in v["items"]], cols=3, fill=TINT_GRAY, label="PROGRAM")
        elif key == "info":
            diagram_title(v["title"])
            kv_table(v["rows"])
            diagram_title("家族会 当日の流れ")
            steps([
                {"title": "受付・送迎", "note": "JR相模原駅 北口より専用送迎車（12:45発・13:00発）"},
                {"title": "開会（13:30）", "note": "はじめての方も、どうぞ安心してお越しください"},
                {"title": "エキスパート講演会", "note": "医師・専門家から、依存症と回復を学ぶ"},
                {"title": "家族ミーティング", "note": "「言いっぱなし・聞きっぱなし」で分かち合う"},
                {"title": "当事者スタッフ面談", "note": "ご家庭の悩みに、回復者スタッフが一緒に向き合う"},
                {"title": "閉会（17:00）", "note": "お帰りも送迎いたします"},
            ])
            caption_line("※ 内容は回により変わることがあります。見学だけの参加も歓迎です。")
            add_image(data["images"]["facility"], caption="家族会の会場 ── 相模原ダルク デイケアセンター（相模原市）。JR相模原駅から送迎車もご用意しています。")
            if v.get("note"):
                callout(v["note"], fill=TINT_BLUE, accent=BLUE)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 私たちが伴走する理由
wr = data["walkReason"]
chapter_band(wr["label"], wr["title"], wr["headline"])
body_lines(wr["body"])
color_box(wr["philosophyLabel"], wr["philosophy"], [], TINT_GREEN, GREEN)
body_lines(wr["body2"])
for l in wr["closing"]:
    p = para(space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, l, bold=True, size=12, color=NAVY)
p = para(space_before=4, align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, wr["tagline"], bold=True, size=18, color=NAVY)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, wr["taglineSub"], bold=True, size=10, color=WARM)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 約束
pr = data["promise"]
chapter_band(pr["chapterLabel"], pr["title"], "")
add_image(data["images"]["walking"])
body_lines(pr["paragraphs"])
cards([{"title": p["title"], "body": p["note"]} for p in pr["pledges"]], cols=3, fill=TINT_GREEN, label=None)
p = para(space_before=6, align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, pr["finalMessage"], bold=True, size=12, color=NAVY)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, pr["signoff"], bold=True, size=11, color=NAVY)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 裏表紙・お問い合わせ
bc = data["backCover"]
chapter_band("CONTACT", "お問い合わせ", bc["headline"])
box = one_cell_box(TINT_BLUE)
cell_para(box, bc["hotline"]["label"], first=True, size=10, bold=True, color=NAVY, space_after=1)
cell_para(box, bc["hotline"]["value"], size=18, bold=True, color=RGBColor(0xC0, 0x5C, 0x5C), space_after=4)
cell_para(box, bc["rep"]["label"], size=10, bold=True, color=NAVY, space_after=1)
cell_para(box, bc["rep"]["value"], size=15, bold=True, color=NAVY, space_after=2)
cell_para(box, bc["support"], size=9.5, space_after=0)
spacer(6)
color_box(bc["eligibility"]["title"], "", [bc["eligibility"]["lead"]] , TINT_GREEN, GREEN)
bullets(bc["eligibility"]["items"], mark="✓", color=GREEN)
para(f'{bc["siteLabel"]}：{bc["site"]}', size=9.5, color=BLUE)
para(data["org"]["name"], bold=True, color=NAVY)
para(bc["disclaimer"], size=8, color=GRAY)

doc.save(OUT)
print(f"Rich DOCX generated: {OUT}")
